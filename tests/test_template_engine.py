"""Unit and integration tests for Track D `src/tools/template_engine.py` (v3.0)."""

from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

from src.tools.template_engine import (
    BUILT_IN_TEMPLATES,
    BUILTIN_TEMPLATES,
    DMBC_HVSC_TEMPLATES,
    NANO_BUILT_IN_BACKDROPS_CATALOG,
    create_presentation_from_template,
    decorate_slide_with_template_chrome,
    get_template_profile,
    list_nano_builtin_backdrops,
    load_template_catalog,
    resolve_default_cover_backdrop_for_badge,
    resolve_effective_slide_style,
)

REPO_ROOT = Path(__file__).resolve().parent.parent


def test_builtin_templates_expanded_to_23_and_29_total() -> None:
    """Verifies 23 built-in templates and 6 DMBC/HVSC templates (29 total built-in catalog items)."""
    assert len(BUILTIN_TEMPLATES) == 23
    assert len(BUILT_IN_TEMPLATES) == 23
    assert len(DMBC_HVSC_TEMPLATES) == 6

    builtin_ids = {t.template_id for t in BUILTIN_TEMPLATES}
    expected_new_ids = {
        "stem_nova_blueprint",
        "eagle_court_of_honor",
        "sea_scouts_maritime",
        "oa_ceremonial_lodge",
        "venturing_alpine_emerald",
        "wood_badge_gilwell",
    }
    assert expected_new_ids.issubset(builtin_ids)

    catalog = load_template_catalog()
    catalog_builtin_ids = {t.template_id for t in catalog if t.is_builtin}
    assert len(catalog_builtin_ids) == 29
    assert expected_new_ids.issubset(catalog_builtin_ids)


def test_nano_builtin_backdrops_expanded_to_24_and_files_exist() -> None:
    """Verifies all 24 photorealistic washed cover backdrops are registered and exist on disk."""
    backdrops = list_nano_builtin_backdrops()
    assert len(NANO_BUILT_IN_BACKDROPS_CATALOG) == 24
    assert len(backdrops) == 24

    backdrop_ids = {b["id"] for b in backdrops}
    expected_new_backdrops = {
        "citizenship_capitol_civic",
        "personal_management_ledger",
        "fitness_cycling_trail",
        "communications_studio",
        "trades_welding_workshop",
        "wildlife_forestry_conservation",
        "astronomy_deep_space_observatory",
        "emergency_medical_triage",
    }
    assert expected_new_backdrops.issubset(backdrop_ids)

    for b in backdrops:
        rel_path = b["path"].lstrip("/")
        disk_path = REPO_ROOT / rel_path
        assert disk_path.exists(), f"Missing backdrop file on disk: {disk_path}"
        assert disk_path.stat().st_size > 1024


def test_resolve_default_cover_backdrop_for_badge_categories() -> None:
    """Verifies badge and category auto-matching for washed cover backdrops."""
    assert "nano_base_citizenship_capitol_civic.jpg" in resolve_default_cover_backdrop_for_badge(
        "Citizenship in the Nation"
    )
    assert "nano_base_personal_management_ledger.jpg" in resolve_default_cover_backdrop_for_badge(
        "Personal Management"
    )
    assert "nano_base_fitness_cycling_trail.jpg" in resolve_default_cover_backdrop_for_badge(
        "Personal Fitness"
    )
    assert "nano_base_communications_studio.jpg" in resolve_default_cover_backdrop_for_badge(
        "Communication"
    )
    assert "nano_base_trades_welding_workshop.jpg" in resolve_default_cover_backdrop_for_badge(
        "Welding"
    )
    assert "nano_base_wildlife_forestry_conservation.jpg" in resolve_default_cover_backdrop_for_badge(
        "Forestry"
    )
    assert "nano_base_astronomy_deep_space_observatory.jpg" in resolve_default_cover_backdrop_for_badge(
        "Astronomy"
    )
    assert "nano_base_emergency_medical_triage.jpg" in resolve_default_cover_backdrop_for_badge(
        "Emergency Preparedness"
    )


def test_auto_cover_backdrop_for_builtin_templates_and_forest_sunrise_cover() -> None:
    """Verifies built-in templates with empty cover_bg_url auto-resolve a badge-matched backdrop on Slide 1."""
    prof = get_template_profile("bsa_scouting_america", badge_name="Welding")
    assert "nano_base_trades_welding_workshop.jpg" in prof.cover_bg_url

    style = resolve_effective_slide_style(
        "bsa_scouting_america",
        deck_style={"cover_backdrop_id": "AUTO_TEMPLATE"},
        badge_name="Astronomy",
    )
    assert "nano_base_astronomy_deep_space_observatory.jpg" in style["cover_bg_url"]

    style_none = resolve_effective_slide_style(
        "bsa_scouting_america",
        deck_style={"cover_backdrop_id": "NONE"},
        badge_name="Astronomy",
    )
    assert style_none["cover_bg_url"] == ""

    # Verify scouting_america_forest_sunrise renders TemplateChrome_CoverBg on Slide 1
    prs, add_slide_fn, _, _ = create_presentation_from_template(
        "scouting_america_forest_sunrise",
        badge_name="Camping",
    )
    cover_slide = add_slide_fn(True, "BEAUTIFIED")
    shape_names = [sh.name for sh in cover_slide.shapes]
    assert "TemplateChrome_CoverBg" in shape_names

    # Verify STUDIO mode renders dark-adapted watermark when canvas_watermark_url is present
    prs2 = Presentation()
    prs2.slide_width = Inches(13.333)
    prs2.slide_height = Inches(7.5)
    blank_slide = prs2.slides.add_slide(prs2.slide_layouts[6])
    stem_prof = get_template_profile("stem_nova_blueprint", badge_name="Robotics")
    decorate_slide_with_template_chrome(
        blank_slide,
        stem_prof,
        is_cover=False,
        tier="STUDIO",
        badge_name="Robotics",
    )
    studio_shape_names = [sh.name for sh in blank_slide.shapes]
    assert "TemplateChrome_Watermark" in studio_shape_names
